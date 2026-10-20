"""
controllers/cosecha_controller.py
---------------------------------
Operational harvest yield (capachos, piece rate, daily floor).

Routes:
  GET /cosecha/detalle                      → page
  GET /api/cosecha/detalle/filters          → cuartel, cuadrilla, cultivo
  GET /api/cosecha/detalle                  → summary + worker rows
  GET /api/cosecha/detalle/download-excel   → xlsx
"""

import io
import logging
import re

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Font

from auth import require_auth
from db import get_connection
from services.cosecha_rendimiento import filter_options, query_detalle

logger = logging.getLogger("controllers.cosecha")

router = APIRouter(dependencies=[Depends(require_auth)])

_templates: Jinja2Templates = None
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def init(templates: Jinja2Templates) -> None:
    global _templates
    _templates = templates


def _blank(value: str | None) -> str | None:
    text = (value or "").strip()
    return text or None


def _cuadrilla(value: str | None) -> int | None:
    text = _blank(value)
    if text is None:
        return None
    if not text.isdigit():
        raise HTTPException(status_code=400, detail="Cuadrilla must be an integer")
    return int(text)


def _dates(fecha_inicio: str, fecha_termino: str) -> tuple[str, str]:
    if not _DATE_RE.match(fecha_inicio or "") or not _DATE_RE.match(fecha_termino or ""):
        raise HTTPException(status_code=400, detail="Dates must be YYYY-MM-DD")
    if fecha_inicio > fecha_termino:
        raise HTTPException(status_code=400, detail="fecha_inicio is after fecha_termino")
    return fecha_inicio, fecha_termino


def _load(
    fecha_inicio: str,
    fecha_termino: str,
    empresa: str | None,
    contratista: str | None,
    cuartel: str | None,
    cuadrilla: str | None,
    cultivo: str | None,
) -> dict:
    inicio, termino = _dates(fecha_inicio, fecha_termino)
    try:
        conn = get_connection()
    except Exception:
        logger.exception("cosecha detalle connection failed")
        raise HTTPException(status_code=503, detail="Error de conexión a la base de datos")
    try:
        with conn.cursor() as cur:
            return query_detalle(
                cur,
                inicio,
                termino,
                empresa=_blank(empresa),
                contratista=_blank(contratista),
                cuartel=_blank(cuartel),
                cuadrilla=_cuadrilla(cuadrilla),
                cultivo=_blank(cultivo),
            )
    finally:
        conn.close()


@router.get("/cosecha/detalle", response_class=HTMLResponse)
async def cosecha_detalle_page(request: Request):
    return _templates.TemplateResponse(request, "cosecha_detalle.html")


@router.get("/api/cosecha/detalle/filters")
async def cosecha_detalle_filters(
    empresa: str = Query(None),
    contratista: str = Query(None),
    fecha_inicio: str = Query(None),
    fecha_termino: str = Query(None),
):
    if fecha_inicio or fecha_termino:
        _dates(fecha_inicio or "", fecha_termino or "")
    try:
        conn = get_connection()
    except Exception:
        logger.exception("cosecha filters connection failed")
        raise HTTPException(status_code=503, detail="Error de conexión a la base de datos")
    try:
        with conn.cursor() as cur:
            return filter_options(
                cur,
                _blank(empresa),
                _blank(contratista),
                _blank(fecha_inicio),
                _blank(fecha_termino),
            )
    finally:
        conn.close()


@router.get("/api/cosecha/detalle")
async def cosecha_detalle_data(
    fecha_inicio: str = Query(...),
    fecha_termino: str = Query(...),
    empresa: str = Query(None),
    contratista: str = Query(None),
    cuartel: str = Query(None),
    cuadrilla: str = Query(None),
    cultivo: str = Query(None),
):
    return _load(
        fecha_inicio, fecha_termino, empresa, contratista, cuartel, cuadrilla, cultivo
    )


def _xlsx(data: dict) -> bytes:
    wb = Workbook()
    resumen = wb.active
    resumen.title = "Rendimiento"
    detalle = wb.create_sheet("Por trabajador")
    header = Font(bold=True)

    resumen_headers = [
        "Fecha",
        "Predio",
        "Contratista",
        "Cuadrilla",
        "Cuartel",
        "Cultivo",
        "Variedad",
        "Trabajadores",
        "Capachos",
        "Capachos por persona",
        "Bins",
        "Trato",
        "Trato + piso",
    ]
    detalle_headers = [
        "Fecha",
        "Predio",
        "Contratista",
        "Cuadrilla",
        "Cuartel",
        "Cultivo",
        "Variedad",
        "Trabajador",
        "Capachos",
        "Precio",
        "Trato",
        "Piso",
        "Trato + piso",
    ]
    for col, label in enumerate(resumen_headers, 1):
        cell = resumen.cell(1, col, label)
        cell.font = header
    for col, label in enumerate(detalle_headers, 1):
        cell = detalle.cell(1, col, label)
        cell.font = header

    money = "#,##0"
    for i, row in enumerate(data["resumen"], 2):
        values = [
            _slash(row["fecha"]),
            row["predio"],
            row["contratista"],
            row["cuadrilla"],
            row["cuartel"],
            row["cultivo"],
            row["variedad"],
            row["trabajadores"],
            row["capachos"],
            row["capachos_persona"],
            row["bins"],
            row["trato"],
            row["a_pagar"],
        ]
        for col, value in enumerate(values, 1):
            cell = resumen.cell(i, col, value)
            if col in (12, 13):
                cell.number_format = money

    for i, row in enumerate(data["rows"], 2):
        values = [
            _slash(row["fecha"]),
            row["predio"],
            row["contratista"],
            row["cuadrilla"],
            row["cuartel"],
            row["cultivo"],
            row["variedad"],
            row["trabajador"],
            row["capachos"],
            row["precio"],
            row["trato"],
            row["piso"],
            row["a_pagar"],
        ]
        for col, value in enumerate(values, 1):
            cell = detalle.cell(i, col, value)
            if col in (10, 11, 12, 13):
                cell.number_format = money

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _slash(iso: str | None) -> str:
    if not iso or len(iso) < 10:
        return iso or ""
    y, m, d = iso[:10].split("-")
    return f"{d}/{m}/{y}"


@router.get("/api/cosecha/detalle/download-excel")
async def cosecha_detalle_excel(
    fecha_inicio: str = Query(...),
    fecha_termino: str = Query(...),
    empresa: str = Query(None),
    contratista: str = Query(None),
    cuartel: str = Query(None),
    cuadrilla: str = Query(None),
    cultivo: str = Query(None),
):
    data = _load(
        fecha_inicio, fecha_termino, empresa, contratista, cuartel, cuadrilla, cultivo
    )
    payload = _xlsx(data)
    filename = f"cosecha-detalle-{fecha_inicio}-{fecha_termino}.xlsx"
    return StreamingResponse(
        io.BytesIO(payload),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
